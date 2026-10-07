import io,json

def make_pptx(title,slides):
    from pptx import Presentation
    prs=Presentation()
    for i,data in enumerate(slides or []):
        slide=prs.slides.add_slide(prs.slide_layouts[0 if i==0 else 1])
        slide.shapes.title.text=data.get("title",title)
        if i:
            tf=slide.placeholders[1].text_frame;tf.clear()
            for j,b in enumerate(data.get("bullets",[])):
                p=tf.paragraphs[0] if j==0 else tf.add_paragraph();p.text=str(b)
    out=io.BytesIO();prs.save(out);return out.getvalue()

def body_from_generated(kind,data):
    if not data:return ""
    if kind=="notes":
        return "\n\n".join(f"## {s.get('heading','Section')}\n"+"\n".join("- "+p for p in s.get("points",[])) for s in data.get("sections",[]))
    if kind=="assignment":
        return "# "+data.get("title","Assignment")+"\n\n"+"\n".join("- "+x for x in data.get("instructions",[]))+"\n\n"+"\n".join(f"{i+1}. {q}" for i,q in enumerate(data.get("questions",[])))
    if kind=="revision":
        return "# "+data.get("title","Revision")+"\n\nKey points:\n"+"\n".join("- "+x for x in data.get("key_points",[]))+"\n\nImportant questions:\n"+"\n".join(f"{i+1}. {x}" for i,x in enumerate(data.get("important_questions",[])))
    if kind in {"question_bank","quiz","ppt"}:return json.dumps(data,indent=2)
    return json.dumps(data,indent=2)
