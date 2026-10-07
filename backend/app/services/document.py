import io,re

def normalize_text(text:str)->str:
    return re.sub(r'\n{3,}','\n\n',re.sub(r'[ \t]+',' ',text)).strip()

def chunk_text(text:str,size:int=1200,overlap:int=180)->list[str]:
    text=normalize_text(text)
    if not text:return []
    out=[];start=0
    while start<len(text):
        end=min(len(text),start+size)
        if end<len(text):
            cut=text.rfind(' ',start,end)
            if cut>start+size//2:end=cut
        out.append(text[start:end].strip())
        if end>=len(text):break
        start=max(end-overlap,start+1)
    return out

def extract_sections(filename:str,data:bytes)->tuple[list[dict],int]:
    ext=filename.lower().rsplit('.',1)[-1] if '.' in filename else ''
    if ext in {'txt','md','csv','json','py','java','js','ts','html','css'}:
        return [{"text":data.decode('utf-8',errors='ignore'),"page":1,"section":"document"}],1
    if ext=='pdf':
        try:
            from pypdf import PdfReader
            pages=PdfReader(io.BytesIO(data)).pages
            return [{"text":p.extract_text() or "","page":i+1,"section":f"page {i+1}"} for i,p in enumerate(pages)],len(pages)
        except Exception:return [],0
    if ext=='docx':
        try:
            from docx import Document
            d=Document(io.BytesIO(data))
            return [{"text":"\n".join(p.text for p in d.paragraphs),"page":1,"section":"document"}],1
        except Exception:return [],0
    if ext=='pptx':
        try:
            from pptx import Presentation
            p=Presentation(io.BytesIO(data)); rows=[]
            for i,slide in enumerate(p.slides,1):
                rows.append({"text":"\n".join(s.text for s in slide.shapes if hasattr(s,'text') and s.text),"page":i,"section":f"slide {i}"})
            return rows,len(rows)
        except Exception:return [],0
    return [],0

def extract_text(filename:str,data:bytes)->tuple[str,int]:
    sections,pages=extract_sections(filename,data)
    return normalize_text('\n\n'.join(x['text'] for x in sections)),pages
